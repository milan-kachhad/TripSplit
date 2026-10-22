import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F1F5F9")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    
    run = p.add_run(code_text.strip())
    run.font.name = 'Courier New'
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(30, 41, 59)
    
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)

def add_callout(doc, text, title="NOTE"):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    
    title_run = p.add_run(f"[{title}] ")
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(29, 78, 216)
    title_run.font.size = Pt(9.5)
    
    text_run = p.add_run(text)
    text_run.font.size = Pt(9.5)
    text_run.font.color.rgb = RGBColor(30, 41, 59)
    
    doc.add_paragraph()

doc = Document()

# Set page margins
for section in doc.sections:
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

# Styles
PRIMARY_COLOR = RGBColor(30, 58, 138)    # Deep Blue
SECONDARY_COLOR = RGBColor(14, 116, 144) # Cyan/Teal
DARK_TEXT = RGBColor(15, 23, 42)         # Slate 900
MUTED_TEXT = RGBColor(100, 116, 139)     # Slate 500

# Document Title Header
title_p = doc.add_paragraph()
title_p.paragraph_format.space_before = Pt(0)
title_p.paragraph_format.space_after = Pt(4)
title_run = title_p.add_run("TripSplit — Project Changes & Implementation Report")
title_run.font.name = 'Calibri'
title_run.font.size = Pt(24)
title_run.font.bold = True
title_run.font.color.rgb = PRIMARY_COLOR

subtitle_p = doc.add_paragraph()
subtitle_p.paragraph_format.space_before = Pt(0)
subtitle_p.paragraph_format.space_after = Pt(16)
sub_run = subtitle_p.add_run("Comprehensive Technical Documentation of Feature Updates, Code Modifications & Database Adjustments")
sub_run.font.name = 'Calibri'
sub_run.font.size = Pt(12)
sub_run.font.color.rgb = MUTED_TEXT

# Metadata Table
meta_tbl = doc.add_table(rows=4, cols=2)
meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
meta_data = [
    ("Project Name", "TripSplit (Shared Trip Expense Manager)"),
    ("Stack", "React (Vite) + Express.js + MySQL (mysql2/promise)"),
    ("Scope of Work", "Email-derived usernames, strict database member addition, dropdown UI, 2% forex conversion charge"),
    ("Date & Status", "October 2026 | Verified (0 lint errors, build passing)")
]
for idx, (label, val) in enumerate(meta_data):
    row = meta_tbl.rows[idx]
    c1, c2 = row.cells[0], row.cells[1]
    c1.width = Inches(1.8)
    c2.width = Inches(4.7)
    set_cell_background(c1, "F8FAFC")
    set_cell_background(c2, "FFFFFF")
    set_cell_margins(c1, 60, 60, 100, 100)
    set_cell_margins(c2, 60, 60, 100, 100)
    
    p1 = c1.paragraphs[0]
    r1 = p1.add_run(label)
    r1.font.bold = True
    r1.font.size = Pt(9.5)
    r1.font.color.rgb = DARK_TEXT
    
    p2 = c2.paragraphs[0]
    r2 = p2.add_run(val)
    r2.font.size = Pt(9.5)
    r2.font.color.rgb = DARK_TEXT

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# Section 1: Executive Summary
h1 = doc.add_heading(level=1)
h1_run = h1.add_run("1. Executive Summary")
h1_run.font.color.rgb = PRIMARY_COLOR

p_summary = doc.add_paragraph(
    "This document provides complete, end-to-end technical documentation of all modifications performed on the TripSplit project. "
    "The updates address four key requirements requested by the project owner:\n"
)
bullet_points = [
    ("Automated Valid Username Derivation: ", "Eliminated manual username errors and formatting bugs by extracting the local username part automatically from the email address before the '@' symbol (e.g. sarthak@example.com -> sarthak) across frontend inputs, context, and backend APIs."),
    ("Strict Database Verification for Trip Members: ", "Enforced that only registered users existing in the MySQL database can be added to a trip. Removed legacy automatic generation of temporary/mock users, ensuring data consistency and referential integrity."),
    ("Clean, Minimalist Dropdown UI for Adding Members: ", "Refactored the Add Member modal from card-based lists and manual input fields into a clean, single native dropdown select list. Removed '@' symbols before usernames and marked existing trip members as disabled to prevent duplicates."),
    ("2.0% Forex Conversion Charges: ", "Integrated a standard 2.0% currency conversion fee whenever foreign currencies (USD, EUR, GBP, THB, AED) are logged and standardized into INR. Updated the database schema, financial calculation engine, expense modal warning, and expense breakdown logs.")
]

for title, desc in bullet_points:
    p_b = doc.add_paragraph(style='List Bullet')
    p_b.paragraph_format.space_after = Pt(4)
    r_t = p_b.add_run(title)
    r_t.font.bold = True
    r_t.font.color.rgb = DARK_TEXT
    r_d = p_b.add_run(desc)
    r_d.font.color.rgb = DARK_TEXT

# Section 2: Summary of Modified & Created Files
h2 = doc.add_heading(level=1)
h2_run = h2.add_run("2. Inventory of Modified & Created Files")
h2_run.font.color.rgb = PRIMARY_COLOR

file_table_data = [
    ("src/pages/Register.jsx", "Modified", "Auto-derive username from email prefix before @, reordered email input first, live @username preview, strict length validation (2-50 chars)."),
    ("src/context/AuthContext.jsx", "Modified", "Updated register() function to automatically fallback to email prefix if username is omitted or blank."),
    ("server/controllers/authController.js", "Modified", "Backend registration extracts local email prefix before @, strips invalid characters, validates [a-zA-Z0-9_.-], and auto-disambiguates collisions."),
    ("server/controllers/userController.js", "Created", "New controller providing GET /api/users to return list of registered users (id, name, username, email) with optional search."),
    ("server/routes/users.js", "Created", "New Express router exposing GET /api/users protected by authMiddleware."),
    ("server/server.js", "Modified", "Mounted /api/users route in the main Express application."),
    ("server/controllers/tripController.js", "Modified", "Strictly verified existing users when adding trip members. Removed mock user auto-creation. Removed unused bcrypt import."),
    ("src/components/AddMemberModal.jsx", "Modified", "Refactored into a single clean <select> dropdown. Removed @ prefix, removed manual inputs, added 'Select User' placeholder."),
    ("src/pages/TripDetails.jsx", "Modified", "Passed existingMembers prop to AddMemberModal; updated expenses line items to display 2% conversion charge breakdown."),
    ("server/utils/currency.js", "Modified", "Added CONVERSION_FEE_PERCENT = 2.0. Computes base converted INR, 2% fee, and total amount."),
    ("server/config/db.js", "Modified", "Added auto-migration check to ensure conversion_fee DECIMAL(10, 2) DEFAULT 0.00 column exists in expenses table."),
    ("server/schema.sql", "Modified", "Added conversion_fee column to expenses table schema definition."),
    ("server/controllers/expenseController.js", "Modified", "Recorded conversion_fee in expenses table during MySQL transaction and returned it in getExpenses."),
    ("src/components/AddExpenseModal.jsx", "Modified", "Added +2% Fee badge and clear explanatory message when a foreign currency is selected."),
    ("PROJECT_DOCUMENTATION.md", "Modified", "Updated architecture notes, Q&A section, member invitation instructions, and multi-currency calculation details.")
]

tbl_files = doc.add_table(rows=len(file_table_data) + 1, cols=3)
tbl_files.alignment = WD_TABLE_ALIGNMENT.CENTER
tbl_files.rows[0].cells[0].width = Inches(2.2)
tbl_files.rows[0].cells[1].width = Inches(1.0)
tbl_files.rows[0].cells[2].width = Inches(3.3)

headers = ["File Path", "Status", "Summary of Changes"]
for c_idx, h_text in enumerate(headers):
    c = tbl_files.rows[0].cells[c_idx]
    set_cell_background(c, "1E3A8A")
    set_cell_margins(c, 80, 80, 100, 100)
    p = c.paragraphs[0]
    r = p.add_run(h_text)
    r.font.bold = True
    r.font.color.rgb = RGBColor(255, 255, 255)
    r.font.size = Pt(9.5)

for r_idx, (fpath, status, fdesc) in enumerate(file_table_data, start=1):
    row = tbl_files.rows[r_idx]
    for c_idx, text in enumerate([fpath, status, fdesc]):
        c = row.cells[c_idx]
        set_cell_background(c, "F8FAFC" if r_idx % 2 == 1 else "FFFFFF")
        set_cell_margins(c, 60, 60, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(text)
        r.font.size = Pt(8.5)
        if c_idx == 1:
            r.font.bold = True
            r.font.color.rgb = RGBColor(16, 185, 129) if status == "Created" else RGBColor(59, 130, 246)
        else:
            r.font.color.rgb = DARK_TEXT

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# Section 3: Deep-Dive Code Explanations
h3 = doc.add_heading(level=1)
h3_run = h3.add_run("3. Detailed Code Modifications by Feature")
h3_run.font.color.rgb = PRIMARY_COLOR

# Feature 1
doc.add_heading(level=2, text="3.1 Feature 1: Email-Derived Username Generation")
doc.add_paragraph(
    "Previously, users were asked to type a username manually, leading to validation errors when spaces or invalid characters "
    "were input. The new implementation automatically derives the username from the email address before the '@' sign, "
    "stripping illegal characters and ensuring real-time UI synchronicity."
)

doc.add_paragraph("Key Code: src/pages/Register.jsx")
add_code_block(doc, """// Helper to extract clean valid username from email prefix before @
const deriveUsernameFromEmail = (emailStr) => {
  if (!emailStr) return '';
  const prefix = emailStr.includes('@') ? emailStr.split('@')[0] : emailStr;
  return prefix.toLowerCase().replace(/[^a-z0-9_.-]/g, '');
};

const handleEmailChange = (e) => {
  const newEmail = e.target.value;
  setEmail(newEmail);
  const autoUsername = deriveUsernameFromEmail(newEmail);
  setUsername(autoUsername);
};""")

doc.add_paragraph("Backend Controller: server/controllers/authController.js")
add_code_block(doc, """// Derive username from email context before '@'
const emailPrefix = trimmedEmail.includes('@')
  ? trimmedEmail.split('@')[0]
  : trimmedEmail;
const cleanEmailPrefix = emailPrefix.toLowerCase().replace(/[^a-z0-9_.-]/g, '');

let finalUsername = (username ? String(username).trim() : cleanEmailPrefix)
  .toLowerCase()
  .replace(/\\s+/g, '')
  .replace(/[^a-z0-9_.-]/g, '');

// Verify username uniqueness - auto-disambiguate with numeric suffix if taken
const [userMatches] = await pool.query(
  'SELECT id FROM users WHERE LOWER(username) = LOWER(?)',
  [finalUsername]
);
if (userMatches.length > 0) {
  finalUsername = `${finalUsername}_${Math.floor(100 + Math.random() * 900)}`;
}""")

# Feature 2
doc.add_heading(level=2, text="3.2 Feature 2: Strict Database User Verification for Trip Members")
doc.add_paragraph(
    "To enforce data integrity, adding members to a trip now strictly requires the user to exist in the database users table. "
    "Legacy auto-provisioning of fake accounts was completely removed. Non-existent users receive a 404 error."
)

doc.add_paragraph("Updated Member Insertion Logic: server/controllers/tripController.js")
add_code_block(doc, """export const addTripMember = async (req, res) => {
  try {
    const tripId = req.params.id;
    const { email, name, userId } = req.body;
    let user = null;

    // 1. Look up by explicit userId if supplied
    if (userId) {
      const [userById] = await pool.query(
        'SELECT id, name, username, email FROM users WHERE id = ?',
        [userId]
      );
      if (userById.length > 0) user = userById[0];
    }

    // 2. Look up user by email or username if not already found
    const input = (email || name || '').trim();
    if (!user && input) {
      const [userRows] = await pool.query(
        'SELECT id, name, username, email FROM users WHERE LOWER(email) = ? OR LOWER(username) = ? OR LOWER(name) = ?',
        [input.toLowerCase(), input.toLowerCase(), input.toLowerCase()]
      );
      if (userRows.length > 0) user = userRows[0];
    }

    // 3. STRICT CHECK: User MUST exist in database
    if (!user) {
      return res.status(404).json({
        message: 'User does not exist in the database. Only registered users can be added as trip members.',
      });
    }

    // 4. Duplicate membership check
    const [existingMembers] = await pool.query(
      'SELECT id FROM trip_members WHERE trip_id = ? AND (user_id = ? OR LOWER(name) = ?)',
      [tripId, user.id, (user.name || user.username).toLowerCase()]
    );
    if (existingMembers.length > 0) {
      return res.status(400).json({ message: 'User is already a member of this trip.' });
    }

    const memberId = 'mem_' + Date.now() + '_' + crypto.randomBytes(3).toString('hex');
    await pool.query(
      'INSERT INTO trip_members (id, trip_id, user_id, name) VALUES (?, ?, ?, ?)',
      [memberId, tripId, user.id, user.username || user.name]
    );
    return res.status(201).json({ message: 'Member added successfully.' });
  } ...
};""")

# Feature 3
doc.add_heading(level=2, text="3.3 Feature 3: Simplified Dropdown Selection UI")
doc.add_paragraph(
    "In accordance with user feedback, the complex multi-card member picker and manual text entry was replaced with "
    "a clean, compact <select> dropdown matching the rest of the application. The '@' symbol was removed before usernames."
)

doc.add_paragraph("Dropdown Implementation: src/components/AddMemberModal.jsx")
add_code_block(doc, """<select
  value={selectedUserId}
  onChange={(e) => setSelectedUserId(e.target.value)}
  disabled={loadingUsers}
  className="w-full px-3 py-2 border border-gray-300 dark:border-slate-700 bg-white dark:bg-slate-800 rounded-md text-sm text-gray-900 dark:text-white focus:outline-none focus:ring-2 focus:ring-blue-500"
>
  <option value="">
    {loadingUsers ? 'Loading users...' : ' Select User '}
  </option>
  {users.map((u) => {
    const alreadyAdded = isMember(u);
    const isCurrent = currentUser && String(currentUser.id) === String(u.id);
    const displayName = u.username || u.name;

    return (
      <option
        key={u.id}
        value={u.id}
        disabled={alreadyAdded}
        className="bg-white dark:bg-slate-800 text-gray-900 dark:text-white"
      >
        {displayName} ({u.email}){isCurrent ? ' (You)' : ''}{alreadyAdded ? ' — Already Added' : ''}
      </option>
    );
  })}
</select>""")

# Feature 4
doc.add_heading(level=2, text="3.4 Feature 4: 2.0% Forex Conversion Charges")
doc.add_paragraph(
    "When members pay bills in foreign currencies (e.g., USD, EUR, GBP), financial institutions apply currency conversion fees. "
    "TripSplit now applies a standard 2.0% forex conversion markup when standardizing foreign transactions into INR."
)

doc.add_paragraph("Mathematical Calculation:")
add_code_block(doc, """Base Converted INR = Foreign Amount * Exchange Rate
Conversion Charge (2%) = Base Converted INR * 0.02
Total Standardized INR = Base Converted INR + Conversion Charge""")

doc.add_paragraph("Currency Conversion Utility: server/utils/currency.js")
add_code_block(doc, """export const CONVERSION_FEE_PERCENT = 2.0; // 2% forex conversion markup fee

export async function convertCurrency(amount, fromCurrency = 'INR', toCurrency = 'INR') {
  const numericAmount = parseFloat(amount) || 0;
  const from = (fromCurrency || 'INR').trim().toUpperCase();
  const to = (toCurrency || 'INR').trim().toUpperCase();

  if (from === to || from === 'INR') {
    return {
      convertedAmount: numericAmount,
      baseConvertedAmount: numericAmount,
      conversionFee: 0,
      feePercentage: 0,
      exchangeRate: 1.0,
      currency: from,
    };
  }

  // Rate fetched via live API or reliable offline fallback
  let rate = await fetchLiveRate(from, to) || FALLBACK_RATES_TO_INR[from] || 1.0;

  const baseConvertedAmount = Math.round(numericAmount * rate * 100) / 100;
  const conversionFee = Math.round(baseConvertedAmount * (CONVERSION_FEE_PERCENT / 100) * 100) / 100;
  const finalTotalAmount = Math.round((baseConvertedAmount + conversionFee) * 100) / 100;

  return {
    convertedAmount: finalTotalAmount,
    baseConvertedAmount,
    conversionFee,
    feePercentage: CONVERSION_FEE_PERCENT,
    exchangeRate: Math.round(rate * 10000) / 10000,
    currency: from,
  };
}""")

# Section 4: Database Schema Updates
h4 = doc.add_heading(level=1)
h4_run = h4.add_run("4. Database Schema Updates & Migrations")
h4_run.font.color.rgb = PRIMARY_COLOR

doc.add_paragraph(
    "The MySQL schema was updated to store the conversion fee for auditing and transparency. "
    "An automatic schema migration check was added to server/config/db.js to ensure the column is provisioned on server startup."
)

doc.add_paragraph("Updated Expenses Table Schema (server/schema.sql):")
add_code_block(doc, """CREATE TABLE expenses (
  id VARCHAR(100) PRIMARY KEY,
  trip_id VARCHAR(100) NOT NULL,
  user_id VARCHAR(100),
  paid_by VARCHAR(100) NOT NULL,
  amount DECIMAL(10, 2) NOT NULL,          -- Final standardized INR total (includes 2% fee)
  description VARCHAR(255) NOT NULL,
  currency VARCHAR(3) DEFAULT 'INR',       -- Foreign currency code (USD, EUR, GBP, etc.)
  original_amount DECIMAL(10, 2),          -- Amount in foreign currency before conversion
  exchange_rate DECIMAL(10, 4) DEFAULT 1.0000,
  conversion_fee DECIMAL(10, 2) DEFAULT 0.00, -- Explicit 2% fee in INR
  split_members TEXT,
  created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
  INDEX idx_expenses_trip (trip_id),
  INDEX idx_expenses_paid_by (paid_by),
  CONSTRAINT fk_expenses_trip FOREIGN KEY (trip_id) REFERENCES trips(id) ON DELETE CASCADE,
  CONSTRAINT fk_expenses_paid_by FOREIGN KEY (paid_by) REFERENCES users(id) ON DELETE CASCADE
) ENGINE=InnoDB DEFAULT CHARSET=utf8mb4 COLLATE=utf8mb4_unicode_ci;""")

# Section 5: Verification & Quality Assurance
h5 = doc.add_heading(level=1)
h5_run = h5.add_run("5. Verification & Quality Assurance")
h5_run.font.color.rgb = PRIMARY_COLOR

qa_points = [
    ("Linter Check (oxlint): ", "Executed oxlint across all 27 files in the project. Passed with 0 errors."),
    ("Production Build (vite build): ", "Executed Vite production bundling. All modules transformed cleanly without bundle errors or missing dependencies."),
    ("Database Migration Verification: ", "Successfully executed the migration check via node. Column conversion_fee confirmed in the live database schema."),
    ("Edge Case Testing: ", "Tested registration with short email prefixes, existing username collisions, invalid non-existent member addition rejection, and multi-currency expense addition.")
]

for title, desc in qa_points:
    p_qa = doc.add_paragraph(style='List Bullet')
    p_qa.paragraph_format.space_after = Pt(4)
    r_t = p_qa.add_run(title)
    r_t.font.bold = True
    r_t.font.color.rgb = DARK_TEXT
    r_d = p_qa.add_run(desc)
    r_d.font.color.rgb = DARK_TEXT

# Save file
output_path = "/Users/sarthakkakadiya229icloud.com/Documents/SEM - 5/TripSplit/TripSplit_Project_Changes.docx"
doc.save(output_path)
print(f"Document successfully created at: {output_path}")
